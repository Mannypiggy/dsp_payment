"""测试：DSP+SA 不被覆盖 / 公式保留 / refreshOnLoad（规格 12 节 132-134）。"""
import re
import zipfile

import openpyxl

from conftest import build_july_ps, july_reconciliation
from processors import reconciliation_processor


def test_pivot_region_not_overwritten(july_reconciliation, tmp_path):
    """修改范围只含数据源，不含 DSP+SA 的 Pivot 区域（132）。"""
    ps = build_july_ps()
    _out, _logs, modified = reconciliation_processor.process_dsp_sa_template(
        july_reconciliation, ps, str(tmp_path), use_com=False)
    assert any("DSP数据源" in m for m in modified)
    assert any("已推广广告数据 数据源" in m for m in modified)
    # 不含 DSP+SA 汇总区域
    for m in modified:
        assert "DSP+SA" not in m


def test_formula_region_not_overwritten(july_reconciliation, tmp_path):
    """D37/E37/F37 GETPIVOTDATA 公式不被覆盖（133）。"""
    ps = build_july_ps()
    out, _logs, _mod = reconciliation_processor.process_dsp_sa_template(
        july_reconciliation, ps, str(tmp_path), use_com=False)
    ws = openpyxl.load_workbook(out, data_only=False)["DSP+SA "]
    assert isinstance(ws["D37"].value, str) and ws["D37"].value.startswith("=GETPIVOTDATA")
    assert isinstance(ws["E37"].value, str) and ws["E37"].value.startswith("=GETPIVOTDATA")
    assert ws["F37"].value == "=E37/D37"


def test_refresh_on_load_set_in_fallback(july_reconciliation, tmp_path):
    """openpyxl 回退路径保存后，pivotCacheDefinition 含 refreshOnLoad（134）。"""
    ps = build_july_ps()
    out, _logs, _mod = reconciliation_processor.process_dsp_sa_template(
        july_reconciliation, ps, str(tmp_path), use_com=False)
    z = zipfile.ZipFile(out)
    for name in z.namelist():
        if re.match(r"xl/pivotCache/pivotCacheDefinition\d+\.xml$", name):
            xml = z.read(name).decode("utf-8")
            assert 'refreshOnLoad="1"' in xml

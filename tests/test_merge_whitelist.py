"""测试：DSP花费明细 merge range 完全保留 + 核算维度沿用模板 + 行数不一致报错。

（本轮返工：禁止 unmerge / 重新 merge / 根据数据数量调整合并区域。）
"""
import os
import shutil

import openpyxl
import pytest

from conftest import build_july_ps, july_spend_detail
from models.payment_summary import DSPOrder, PaymentSummary
from processors import spend_detail_processor

EXPECTED_MERGES = {"B3:B13", "C3:C13", "D3:D13", "E3:E13", "F3:F8", "F9:F13", "G3:G13"}


def _merges(path) -> set:
    ws = openpyxl.load_workbook(path)["示例"]
    return {str(m) for m in ws.merged_cells.ranges}


def test_merge_ranges_fully_preserved(july_spend_detail, tmp_path):
    """处理前后 merge range 完全一致（不是只比较数量）。"""
    ps = build_july_ps()
    before = _merges(july_spend_detail)
    assert before == EXPECTED_MERGES
    out, _ = spend_detail_processor.process_spend_detail(july_spend_detail, ps, str(tmp_path))
    after = _merges(out)
    assert after == before == EXPECTED_MERGES


def _make_store_template(july_spend_detail, tmp_path):
    """复制模板并预填 F 列核算维度 / G 列店铺品牌，用于验证不被程序改写。"""
    src = os.path.join(str(tmp_path), "模板_带店铺.xlsx")
    shutil.copy(july_spend_detail, src)
    wb = openpyxl.load_workbook(src)
    ws = wb["示例"]
    ws["F3"] = "Amazon美国站哥贝尔店"
    ws["F9"] = "Amazon美国站DeePaint店"
    ws["G3"] = "HUION"
    ws["C3"] = "美国"
    ws["D3"] = "USD"
    wb.save(src)
    return src


def test_store_dimension_preserved(july_spend_detail, tmp_path):
    """F3=哥贝尔店 / F9=DeePaint店 处理后均不变（F 列沿用模板，不从 Order 推导）。"""
    ps = build_july_ps()
    src = _make_store_template(july_spend_detail, tmp_path)
    out, _ = spend_detail_processor.process_spend_detail(src, ps, str(tmp_path))
    ws = openpyxl.load_workbook(out)["示例"]
    assert ws["F3"].value == "Amazon美国站哥贝尔店"
    assert ws["F9"].value == "Amazon美国站DeePaint店"
    assert ws["G3"].value == "HUION"
    assert ws["C3"].value == "美国"
    assert ws["D3"].value == "USD"


def test_order_count_mismatch_raises(july_spend_detail, tmp_path):
    """Order 数量 != 模板可填写行数 → 报错，不调整 merge。"""
    ps = PaymentSummary(month_code="202608")
    ps.dsp_orders = [DSPOrder(order=f"Campaign_{i}", strategy="Conversion") for i in range(5)]
    with pytest.raises(ValueError, match="模板行数与当月Order数量不一致"):
        spend_detail_processor.process_spend_detail(july_spend_detail, ps, str(tmp_path))


def test_apply_amount_written_to_E3(july_spend_detail, tmp_path):
    """E3 = 对账单结算金额（45103.48）；C/D/F/G 已填写。"""
    ps = build_july_ps()
    out, _ = spend_detail_processor.process_spend_detail(july_spend_detail, ps, str(tmp_path))
    ws = openpyxl.load_workbook(out)["示例"]
    assert float(ws["E3"].value) == pytest.approx(45103.48, abs=0.01)
    assert ws["B3"].value == "202607"
    assert ws["C3"].value == "美国"
    assert ws["D3"].value == "USD"
    assert ws["F3"].value == "Amazon美国站哥贝尔店"
    assert ws["G3"].value == "HUION"


def test_store_column_fills_each_merge(july_spend_detail, tmp_path):
    """F 列上下两个合并区域分别填写不同店铺。"""
    ps = build_july_ps()
    stores = ["Amazon美国站哥贝尔店", "Amazon美国站DeePaint店"]
    out, _ = spend_detail_processor.process_spend_detail(
        july_spend_detail, ps, str(tmp_path), stores=stores)
    ws = openpyxl.load_workbook(out)["示例"]
    assert ws["F3"].value == "Amazon美国站哥贝尔店"
    assert ws["F9"].value == "Amazon美国站DeePaint店"

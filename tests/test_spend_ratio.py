"""测试：店铺花费占比（规格 12 节 119-122 + 129-131）。"""
import os

import openpyxl
import pytest

from conftest import build_july_ps, july_spend_ratio
from processors import spend_ratio_processor
from utils import excel_com

EXPECTED_STORES = ["Amazon美国站哥贝尔店", "Amazon美国站ArtCave店",
                   "Amazon美国站DeePaint店", "Amazon美国站深圳矢原店"]
EXPECTED_RATIOS = {"Amazon美国站哥贝尔店": 0.8, "Amazon美国站ArtCave店": 0.0,
                   "Amazon美国站DeePaint店": 0.2, "Amazon美国站深圳矢原店": 0.0}


def _read_store_rows(path):
    """读取输出文件的店铺行（名字 -> (row, B金额, C比例)）。

    只取 C 列（比例）为数字的行，排除标题/注释。
    """
    import xlrd
    if path.lower().endswith(".xls") and not path.lower().endswith(".xlsx"):
        rb = xlrd.open_workbook(path)
        ws = rb.sheet_by_index(0)
        out = {}
        for r in range(ws.nrows):
            name = str(ws.cell_value(r, 0)).strip()
            ratio = ws.cell_value(r, 2)
            if name and name not in ("核算维度", "合计") and isinstance(ratio, (int, float)):
                out[name] = (r + 1, ws.cell_value(r, 1), ratio)
        return out
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    out = {}
    for r in range(1, ws.max_row + 1):
        name = str(ws.cell(r, 1).value or "").strip()
        ratio = ws.cell(r, 3).value
        if name and name not in ("核算维度", "合计") and isinstance(ratio, (int, float)):
            out[name] = (r, ws.cell(r, 2).value, ratio)
    return out


def test_store_names_preserved(july_spend_ratio, tmp_path):
    ps = build_july_ps()
    out, _ = spend_ratio_processor.process_spend_ratio(july_spend_ratio, ps, str(tmp_path))
    rows = _read_store_rows(out)
    assert set(rows.keys()) == set(EXPECTED_STORES)


def test_store_ratios_preserved(july_spend_ratio, tmp_path):
    ps = build_july_ps()
    out, _ = spend_ratio_processor.process_spend_ratio(july_spend_ratio, ps, str(tmp_path))
    rows = _read_store_rows(out)
    for name, ratio in EXPECTED_RATIOS.items():
        assert abs(float(rows[name][2]) - ratio) < 1e-9, f"{name} 比例被改"


def test_no_order_to_store_derivation():
    """店铺/比例来自模板，不根据 Order 推导。"""
    ps = build_july_ps()
    # PaymentSummary 上不应存在按 Order 的 store 映射
    for o in ps.dsp_orders:
        # 单一默认标签，不是按 Order 推导
        assert o.store == "Amazon美国站哥贝尔店"


def test_amount_calculated_by_fixed_ratio(july_spend_ratio, tmp_path):
    ps = build_july_ps()
    out, _ = spend_ratio_processor.process_spend_ratio(july_spend_ratio, ps, str(tmp_path))
    rows = _read_store_rows(out)
    apply_amount = float(ps.payment_amount)
    assert abs(float(rows["Amazon美国站哥贝尔店"][1]) - apply_amount * 0.8) < 0.01
    assert abs(float(rows["Amazon美国站DeePaint店"][1]) - apply_amount * 0.2) < 0.01
    assert float(rows["Amazon美国站ArtCave店"][1]) == 0.0


def test_com_keeps_xls_when_available(july_spend_ratio, tmp_path):
    """COM 可用时输出保持 .xls（129）。"""
    if not excel_com.com_available():
        pytest.skip("本机无 Excel COM")
    ps = build_july_ps()
    out, logs = spend_ratio_processor.process_spend_ratio(july_spend_ratio, ps, str(tmp_path))
    assert out.lower().endswith(".xls") and not out.lower().endswith(".xlsx")
    assert any("保留 .xls" in l or "COM" in l for l in logs)


def test_com_unavailable_warns_and_falls_back(july_spend_ratio, tmp_path, monkeypatch):
    """COM 不可用时明确警告并回退 .xlsx（130）。"""
    monkeypatch.setattr(excel_com, "com_available", lambda: False)
    ps = build_july_ps()
    out, logs = spend_ratio_processor.process_spend_ratio(july_spend_ratio, ps, str(tmp_path))
    assert out.lower().endswith(".xlsx")
    assert any("无法原格式编辑 .xls" in l or "已转换为 .xlsx" in l for l in logs)


def test_no_silent_xls_to_xlsx(july_spend_ratio, tmp_path, monkeypatch):
    """回退必须产生告警日志，不能静默（131）。"""
    monkeypatch.setattr(excel_com, "com_available", lambda: False)
    ps = build_july_ps()
    _out, logs = spend_ratio_processor.process_spend_ratio(july_spend_ratio, ps, str(tmp_path))
    assert any("无法原格式编辑" in l for l in logs)

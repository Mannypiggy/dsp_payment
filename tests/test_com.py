"""测试：Windows Excel COM 路径（Pivot 完整性 + 公式保留）。

本机有 Excel COM 时验证 COM 路径能正确刷新 Pivot；无则跳过。
"""
import openpyxl
import pytest

from conftest import build_july_ps, july_reconciliation
from processors import reconciliation_processor
from utils import excel_com


def test_com_reconciliation_pivot_refreshed(july_reconciliation, tmp_path):
    if not excel_com.com_available():
        pytest.skip("本机无 Excel COM")
    ps = build_july_ps()
    out, logs, _mod = reconciliation_processor.process_dsp_sa_template(
        july_reconciliation, ps, str(tmp_path), use_com=True)
    assert any("COM" in l for l in logs)

    ws = openpyxl.load_workbook(out, data_only=True)["DSP+SA "]
    # Pivot 刷新后的正确值（与 2026-07 历史一致）
    assert ws["B4"].value == 5294118
    assert ws["B6"].value == 13669440
    assert ws["B25"].value == 16405160
    assert abs(ws["I6"].value - 49025.52) < 0.01
    # GETPIVOTDATA 公式综合行
    assert abs(ws["D37"].value - 194247.92) < 0.01
    assert abs(ws["E37"].value - 965136.46) < 0.01


def test_com_reconciliation_formulas_preserved(july_reconciliation, tmp_path):
    if not excel_com.com_available():
        pytest.skip("本机无 Excel COM")
    ps = build_july_ps()
    out, _logs, _mod = reconciliation_processor.process_dsp_sa_template(
        july_reconciliation, ps, str(tmp_path), use_com=True)
    ws = openpyxl.load_workbook(out, data_only=False)["DSP+SA "]
    assert ws["D37"].value.startswith("=GETPIVOTDATA")
    assert ws["F37"].value == "=E37/D37"

"""上线前收口测试：列宽恢复 / 隐藏状态 / COM 等待与清理 / 配置安全（145-164）。"""
import os

import openpyxl
import pytest

from conftest import JULY_BASE, build_july_ps, july_reconciliation
from processors import reconciliation_processor
from utils import excel_com, excel_utils, validation


# ---- 列宽 / 行高 / 隐藏（145-147）----
class TestComDimensions:
    def test_com_column_width_restored(self, july_reconciliation, tmp_path):
        """COM 刷新后列宽恢复（145）。"""
        if not excel_com.com_available():
            pytest.skip("本机无 Excel COM")
        ps = build_july_ps()
        out, _logs, _mod = reconciliation_processor.process_dsp_sa_template(
            july_reconciliation, ps, str(tmp_path), use_com=True)
        before = excel_utils.workbook_signature(openpyxl.load_workbook(july_reconciliation))
        after = excel_utils.workbook_signature(openpyxl.load_workbook(out))
        res = validation.compare_signatures(before, after)
        assert res.ok, res.differences

    def test_com_hidden_state_preserved(self, july_reconciliation, tmp_path):
        """隐藏行列状态保留（147）。"""
        if not excel_com.com_available():
            pytest.skip("本机无 Excel COM")
        ps = build_july_ps()
        out, _logs, _mod = reconciliation_processor.process_dsp_sa_template(
            july_reconciliation, ps, str(tmp_path), use_com=True)
        b = openpyxl.load_workbook(july_reconciliation)
        a = openpyxl.load_workbook(out)
        for sn in b.sheetnames:
            bh = {k for k, v in b[sn].column_dimensions.items() if v.hidden}
            ah = {k for k, v in a[sn].column_dimensions.items() if v.hidden}
            assert bh == ah, f"{sn} 隐藏列变化"


# ---- COM 等待 / 超时 / 清理（155-157）----
class TestComWait:
    def test_wait_refresh_done(self):
        class App:
            def CalculateUntilAsyncQueriesDone(self):
                return None
        assert excel_com._wait_for_refresh(App(), timeout=1) is True

    def test_wait_refresh_timeout(self):
        class App:
            CalculationState = 1  # 永远在计算中，且无 CalculateUntilAsyncQueriesDone
        assert excel_com._wait_for_refresh(App(), timeout=1) is False

    def test_com_exception_finally_quit(self, monkeypatch):
        """COM 异常时 finally 正确 Quit（157）。"""
        calls = {"quit": 0}

        class MockWorkbooks:
            def Open(self, *a, **k):
                raise RuntimeError("open fail")

        class MockApp:
            def __init__(self):
                self.Workbooks = MockWorkbooks()
            def Quit(self):
                calls["quit"] += 1

        monkeypatch.setattr(excel_com, "_dispatch", lambda: MockApp())
        ok = excel_com.write_reconciliation_via_com("x", "y", [], [])
        assert ok is False
        assert calls["quit"] == 1


# ---- 配置安全（162-164）----
class TestConfigSecurity:
    def _root(self):
        return os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

    def test_gitignore_covers_env_and_local_config(self):
        """.env / config.local.yaml 被 .gitignore（162/163）。"""
        gi = os.path.join(self._root(), ".gitignore")
        assert os.path.exists(gi)
        content = open(gi, encoding="utf-8").read()
        assert ".env" in content
        assert "config.local.yaml" in content
        assert "config.yaml" in content

    def test_example_config_is_desensitized(self):
        """example 配置不包含真实银行账号/邮箱（164）。"""
        p = os.path.join(self._root(), "config.example.yaml")
        assert os.path.exists(p)
        content = open(p, encoding="utf-8").read()
        for secret in ["124-787771-838", "HSBCHKHHHKH", "crystal@huiontablet.com",
                       "songmeimei@huion.cn", "SPARKX MARKETING", "HUION GLOBAL (HK)"]:
            assert secret not in content, f"example 配置泄露敏感信息：{secret}"
        assert "YOUR_" in content

    def test_env_example_is_desensitized(self):
        p = os.path.join(self._root(), ".env.example")
        assert os.path.exists(p)
        content = open(p, encoding="utf-8").read()
        assert "sk-" not in content.lower()

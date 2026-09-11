"""解析已推广广告数据，执行 SB2 转换 + SD 销售额修正，聚合站内指标。"""
from __future__ import annotations

import datetime
import re
from typing import List, Optional, Tuple

import openpyxl

from models.payment_summary import PaymentSummary, PromotedAd, SAChannelSummary
from rules import promoted_ads_mapping
from processors import order_summary_processor
from utils import money


def _to_float(v) -> float:
    if v is None or v == "":
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


_DATE_COL_KEYWORDS = ["日期", "时间", "统计", "月份", "月", "Date", "Month", "Interval", "Period"]


def detect_data_month(path: str) -> Optional[Tuple[int, int]]:
    """从已推广广告源数据识别统计月份（若含日期/月份列）。

    识别不到返回 None（调用方跳过月份校验并提示）。
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [[c for c in row] for row in ws.iter_rows(values_only=True)]
    header_idx = promoted_ads_mapping.find_header_row(rows)
    if header_idx is None:
        return None
    header = [str(h) if h is not None else "" for h in rows[header_idx]]

    months = set()
    for i, h in enumerate(header):
        if not any(k in h for k in _DATE_COL_KEYWORDS):
            continue
        for row in rows[header_idx + 1:]:
            if row is None or i >= len(row):
                continue
            v = row[i]
            if isinstance(v, (datetime.datetime, datetime.date)):
                months.add((v.year, v.month))
            elif isinstance(v, str):
                m = re.search(r"(20\d{2})[/\-.]?(\d{1,2})", v)
                if m:
                    months.add((int(m.group(1)), int(m.group(2))))
    if len(months) == 1:
        return months.pop()
    return None


def parse_promoted_ads(path: str) -> List[PromotedAd]:
    """读取「已推广-汇总」源数据，执行 SB2 转换与 SD 修正。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [[c for c in row] for row in ws.iter_rows(values_only=True)]
    header_idx = promoted_ads_mapping.find_header_row(rows)
    if header_idx is None:
        raise ValueError("已推广广告数据未找到表头行（需含「类型」且含「花费」或「销售额」）")
    header = rows[header_idx]
    fmap = promoted_ads_mapping.build_field_map(header)

    def col(field):
        i = fmap.get(field)
        return row[i] if i is not None else None

    ads: List[PromotedAd] = []
    for row in rows[header_idx + 1:]:
        if row is None or all(c is None or c == "" for c in row):
            continue
        raw_type = str(col("type") or "").strip()
        campaign = str(col("campaign") or "").strip()
        portfolio = str(col("portfolio") or "").strip()
        if not raw_type:
            continue

        ad = PromotedAd(
            ad_type=promoted_ads_mapping.convert_type(raw_type, campaign, portfolio),
            portfolio=portfolio,
            campaign=campaign,
            ad_group=str(col("ad_group") or "").strip(),
            target_type=str(col("target_type") or "").strip(),
            asin=str(col("asin") or "").strip(),
            msku=str(col("msku") or "").strip(),
            status=str(col("status") or "").strip(),
            impressions=_to_float(col("impressions")),
            clicks=_to_float(col("clicks")),
            cost=_to_float(col("cost")),
            sales=_to_float(col("sales")),
            orders=_to_float(col("orders")),
            direct_sales=_to_float(col("direct_sales")),
        )
        promoted_ads_mapping.apply_sd_rule(ad)
        ads.append(ad)
    return ads


def aggregate_onsite(ps: PaymentSummary, ads: List[PromotedAd]) -> None:
    """按类型聚合站内广告，写入 PaymentSummary。"""
    ps.promoted_ads = ads
    channels: List[SAChannelSummary] = []
    for t in ["SB", "SBV", "SD", "SP"]:
        sub = [a for a in ads if a.ad_type == t]
        if not sub:
            continue
        ch = SAChannelSummary(
            ad_type=t,
            impressions=sum(a.impressions for a in sub),
            clicks=sum(a.clicks for a in sub),
            cost=sum(a.cost for a in sub),
            sales=sum(a.sales for a in sub),
            orders=sum(a.orders for a in sub),
        )
        channels.append(ch)

    ps.onsite_channels = channels
    ps.onsite_impressions = sum(ch.impressions for ch in channels)
    ps.onsite_clicks = sum(ch.clicks for ch in channels)
    ps.onsite_orders = sum(ch.orders for ch in channels)
    # 核心财务金额用 Decimal：逐条广告累加（避免先 float 求和引入误差）
    ps.onsite_cost = sum(money.D(a.cost) for a in ads)
    ps.onsite_sales = sum(money.D(a.sales) for a in ads)
    ps.onsite_roas = float(ps.onsite_sales / ps.onsite_cost) if ps.onsite_cost else 0.0

    order_summary_processor.refresh_combined(ps)

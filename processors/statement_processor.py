"""付款对账单处理：读取账单金额，与程序计算值核对。

对账单 = 付款金额核对依据，**不是** DSP数据源 / 已推广广告数据源 的存储容器。
"""
from __future__ import annotations

from dataclasses import dataclass
from typing import List, Optional

import openpyxl

from models.payment_summary import PaymentSummary


@dataclass
class StatementData:
    dsp_cost: Optional[float] = None    # 广告消耗（当地货币）
    rebate: Optional[float] = None       # 返点金额
    payment: Optional[float] = None      # 结算金额 / 账单金额
    rebate_rate: Optional[float] = None  # 返点比例


def read_statement(path: str) -> Optional[StatementData]:
    """读取对账单中的付款金额。识别不到返回 None。"""
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception:
        return None

    # 找对账 Sheet（对账清单 / 待付款）
    ws = None
    for sn in wb.sheetnames:
        if "对账" in sn or "清单" in sn or "待付款" in sn:
            ws = wb[sn]
            break
    if ws is None:
        return None

    header_row = None
    header = []
    for r in range(1, min(ws.max_row, 30) + 1):
        row_vals = [str(ws.cell(r, c).value or "") for c in range(1, ws.max_column + 1)]
        joined = "|".join(row_vals)
        if "广告消耗" in joined and ("返点" in joined or "结算" in joined):
            header_row = r
            header = row_vals
            break
    if header_row is None:
        return None

    def col_idx(keyword: str, exclude: str = None) -> Optional[int]:
        for i, h in enumerate(header):
            if keyword in h and (exclude is None or exclude not in h):
                return i + 1
        return None

    cost_col = col_idx("广告消耗")
    rebate_col = col_idx("返点金额")
    # 「结算金额」要排除「广告结算金额」（那是返点前的结算额）
    payment_col = col_idx("结算金额", exclude="广告")

    def get(col: Optional[int]) -> Optional[float]:
        if col is None:
            return None
        v = ws.cell(header_row + 1, col).value
        return float(v) if isinstance(v, (int, float)) else None

    data = StatementData(
        dsp_cost=get(cost_col),
        rebate=get(rebate_col),
        payment=get(payment_col),
    )
    # 返点比例（0<rate<1 的数值，通常紧邻返点金额）
    if rebate_col is not None:
        for r in range(header_row, header_row + 3):
            for c in range(1, ws.max_column + 1):
                v = ws.cell(r, c).value
                if isinstance(v, (int, float)) and 0 < v < 1:
                    data.rebate_rate = float(v)
    if data.dsp_cost is None and data.rebate is None and data.payment is None:
        return None
    return data


def verify_statement(ps: PaymentSummary, statement: Optional[StatementData]) -> List[dict]:
    """核对程序计算值 vs 对账单读取值。返回 [{"name","status","message"}]。"""
    if statement is None:
        return [{"name": "对账单金额核对", "status": "warn",
                 "message": "无法自动识别对账单付款金额，请人工确认"}]

    checks: List[dict] = []

    if statement.dsp_cost is not None:
        diff = abs(float(ps.billing_base_cost) - statement.dsp_cost)
        ok = diff <= 0.01
        checks.append({"name": "对账单 DSP费用", "status": "pass" if ok else "fail",
                       "message": f"程序 {float(ps.billing_base_cost):,.2f} vs 对账单 {statement.dsp_cost:,.2f}"})

    if statement.rebate is not None:
        diff = abs(float(ps.rebate_raw) - statement.rebate)
        ok = diff <= 0.01
        checks.append({"name": "对账单 返点金额", "status": "pass" if ok else "fail",
                       "message": f"程序 {float(ps.rebate_raw):,.2f} vs 对账单 {statement.rebate:,.2f}"})

    if statement.payment is not None:
        diff = abs(float(ps.payment_raw) - statement.payment)
        ok = diff <= 0.01
        checks.append({"name": "对账单 结算金额", "status": "pass" if ok else "fail",
                       "message": f"程序 {float(ps.payment_raw):,.2f} vs 对账单 {statement.payment:,.2f}"})

    if not checks:
        checks.append({"name": "对账单金额核对", "status": "warn",
                       "message": "对账单中未识别到广告消耗/返点/结算金额，请人工确认"})
    return checks

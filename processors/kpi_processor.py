"""KPI 计算 + 口径核对 +（可选）KPI 模板填值。

单一来源：PaymentSummary.kpi_results；页面 / Excel / 邮件都读这一份。

KPI Excel 只在用户明确上传《KPI模板.xlsx》时才生成——复制模板、原样填值，
**禁止程序自行设计/新建 KPI 表**。未上传模板时只计算/展示/生成原因/用于邮件。
"""
from __future__ import annotations

import os
import shutil
from typing import List, Optional, Tuple

import openpyxl

from models.payment_summary import PaymentSummary
from rules import kpi_rules


def compute_and_reconcile(ps: PaymentSummary, targets: Optional[dict] = None,
                          template_values: Optional[dict] = None,
                          diff_tolerance: float = 0.02) -> List:
    """计算 KPI 并做口径核对（模板值 vs 计算值），结果写回 ps.kpi_results。"""
    results = kpi_rules.compute_kpis(ps, targets=targets, diff_tolerance=diff_tolerance)
    kpi_rules.reconcile_template_kpi(results, template_values, diff_tolerance)
    ps.kpi_results = results
    ps.kpi_targets = {r.name: r.target for r in results}
    # reconcile 会覆盖 r.final，这里同步刷新字典，确保 kpi_final / kpi_actual 与 final 一致
    ps.kpi_calculated = {r.name: r.calculated for r in results}
    ps.kpi_final = {r.name: r.final for r in results}
    ps.kpi_actual = {r.name: r.final for r in results}
    ps.kpi_status = {r.name: ("达标" if r.is_met else "未达标") for r in results}
    return results


def fill_kpi_template(template_path: str, ps: PaymentSummary, output_dir: str,
                      output_name: str = None) -> Tuple[str, List[str]]:
    """复制用户上传的 KPI 模板，原样填入当月 KPI 值。不改结构、不重新设计。

    定位策略（通用，适配常见 KPI 表）：
      - 先在表头行找「实际/实际完成」列、可选「是否达标」「原因」列；
      - 再在数据行中按 KPI 指标名匹配，把 final（最终汇报值）/ 达标状态 / 原因 填入对应列。
    """
    logs: List[str] = []
    if output_name is None:
        output_name = f"DSP_KPI-{ps.month_code}_已处理.xlsx"
    out_path = os.path.join(output_dir, output_name)
    shutil.copy(template_path, out_path)

    wb = openpyxl.load_workbook(out_path)
    by_name = {r.name: r for r in ps.kpi_results}
    filled = 0

    for ws in wb.worksheets:
        # 1) 找表头行：含「实际」或「指标」或「KPI」
        header_row = None
        header_map = {}  # 列名 -> 列号
        for row in ws.iter_rows(min_row=1, max_row=10):
            for cell in row:
                v = str(cell.value or "").strip()
                if "实际" in v or "指标" in v or "KPI" in v or "目标" in v:
                    header_row = cell.row
                    break
            if header_row is not None:
                break
        if header_row is None:
            continue

        for cell in ws[header_row]:
            v = str(cell.value or "").strip()
            if not v:
                continue
            for key, label in (("final", "实际"), ("status", "达标"), ("status", "状态"),
                               ("reason", "原因")):
                if label in v:
                    header_map.setdefault(key, cell.column)
            if "目标" in v:
                header_map.setdefault("target", cell.column)

        # 2) 在数据行里按 KPI 名填值
        name_col = None
        for cell in ws[header_row]:
            v = str(cell.value or "").strip()
            if v in ("KPI指标", "指标", "KPI", "指标名称", "KPI 指标"):
                name_col = cell.column
                break
        if name_col is None:
            continue

        for row in ws.iter_rows(min_row=header_row + 1):
            name_cell = row[name_col - 1]
            name = str(name_cell.value or "").strip()
            if name not in by_name:
                continue
            r = by_name[name]
            if "target" in header_map and row[header_map["target"] - 1].value in (None, ""):
                row[header_map["target"] - 1].value = r.target
            if "final" in header_map and row[header_map["final"] - 1].value in (None, ""):
                row[header_map["final"] - 1].value = r.final
            if "status" in header_map and row[header_map["status"] - 1].value in (None, ""):
                row[header_map["status"] - 1].value = "达标" if r.is_met else "未达标"
            if "reason" in header_map and row[header_map["reason"] - 1].value in (None, ""):
                # 原因列优先写入 kpi_reason_final（已人工确认），其次 AI 提议的 r.reason
                row[header_map["reason"] - 1].value = ps.kpi_reason_final.get(r.name) or r.reason
            filled += 1

    wb.save(out_path)
    if filled == 0:
        logs.append(f"⚠ KPI 模板中未定位到 KPI 指标行，已复制模板原样输出：{output_name}")
    else:
        logs.append(f"✅ KPI 模板已复制并填入 {filled} 个 KPI 值：{output_name}")
    return out_path, logs

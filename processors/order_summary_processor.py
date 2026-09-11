"""解析 Amazon DSP Order Summary → DSPOrder 列表，聚合到 PaymentSummary。"""
from __future__ import annotations

import datetime
from typing import List, Optional, Tuple

import openpyxl

from models.payment_summary import DSPOrder, PaymentSummary, StrategySummary
from rules import dsp_mapping
from utils import money


def _to_float(v) -> float:
    if v is None or v == "":
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def detect_data_month(path: str) -> Optional[Tuple[int, int]]:
    """从 Order Summary 的 Interval start / Interval end 识别数据月份。

    返回 (year, month)；识别不到返回 None。
    """
    interval = detect_interval(path)
    if interval is None:
        return None
    start, end = interval
    if start.year == end.year and start.month == end.month:
        return (start.year, start.month)
    return None


def detect_interval(path: str) -> Optional[Tuple[datetime.datetime, datetime.datetime]]:
    """从 Order Summary 的 Interval start / Interval end 识别统计区间。

    返回 (最早 start, 最晚 end)；识别不到返回 None。
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [[c for c in row] for row in ws.iter_rows(values_only=True)]
    header_idx = dsp_mapping.find_header_row(rows)
    if header_idx is None:
        return None
    header = [str(h) if h is not None else "" for h in rows[header_idx]]

    def col_idx(name):
        for i, h in enumerate(header):
            if h.strip() == name:
                return i
        return None

    start_col = col_idx("Interval start")
    end_col = col_idx("Interval end")
    starts = []
    ends = []
    for row in rows[header_idx + 1:]:
        if row is None:
            continue
        sv = row[start_col] if start_col is not None and start_col < len(row) else None
        ev = row[end_col] if end_col is not None and end_col < len(row) else None
        if isinstance(sv, datetime.datetime):
            starts.append(sv)
        if isinstance(ev, datetime.datetime):
            ends.append(ev)
    if not starts:
        return None
    return (min(starts), max(ends))


def parse_order_summary(path: str, store_default: str = "") -> List[DSPOrder]:
    """读取 Order Summary xlsx，返回 DSPOrder 列表。

    - 字段映射按 dsp_mapping.DSP_FIELD_ALIASES，找不到的字段跳过并记录到 order 之外（由调用方记 warning）。
    - Strategy 由标题识别，识别不到 -> strategy 为空（调用方须报错停止）。
    - 核算维度用单一默认店铺标签（不做 Order→Store 推导）。
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    # 读取所有行
    rows = [[c for c in row] for row in ws.iter_rows(values_only=True)]
    header_idx = dsp_mapping.find_header_row(rows)
    if header_idx is None:
        raise ValueError("Order Summary 未找到表头行（需含 Total cost 与 Impressions）")
    header = rows[header_idx]
    fmap = dsp_mapping.build_field_map(header)

    def col(field):
        i = fmap.get(field)
        return row[i] if i is not None else None

    orders: List[DSPOrder] = []
    for row in rows[header_idx + 1:]:
        if row is None or all(c is None or c == "" for c in row):
            continue
        name = col("order")
        if name is None or str(name).strip() == "":
            continue

        strategy = dsp_mapping.detect_strategy(str(name))
        order = DSPOrder(
            order=str(name).strip(),
            order_id=str(col("order_id") or "").strip(),
            strategy=strategy,
            total_cost=_to_float(col("total_cost")),
            impressions=_to_float(col("impressions")),
            clicks=_to_float(col("clicks")),
            ctr=_to_float(col("ctr")),
            total_dpv=_to_float(col("total_dpv")),
            total_dpvr=_to_float(col("total_dpvr")),
            total_atc=_to_float(col("total_atc")),
            total_atcr=_to_float(col("total_atcr")),
            promoted_sales=_to_float(col("promoted_sales")),
            promoted_roas=_to_float(col("promoted_roas")),
            total_purchases=_to_float(col("total_purchases")),
            total_purchase_rate=_to_float(col("total_purchase_rate")),
            total_units_sold=_to_float(col("total_units_sold")),
            total_sales=_to_float(col("total_sales")),
            total_roas=_to_float(col("total_roas")),
        )
        # 核算维度：单一默认店铺标签（不做 Order→Store 推导）
        order.store = store_default
        orders.append(order)
    return orders


def aggregate(ps: PaymentSummary, orders: List[DSPOrder], rebate_rate: float = 0.08) -> None:
    """把 orders 聚合进 PaymentSummary（DSP 汇总、Strategy 汇总、返点、付款）。

    核心财务金额统一用 Decimal（避免二进制浮点误差）。
    返点口径（与 2026-07 对账清单一致）：
      dsp_cost 先 quantize 到 2 位（49025.52）→ rebate = dsp_cost(2位) × rate（3922.0416）
      → payment = dsp_cost(2位) - rebate_raw（45103.4784），展示再舍 2 位。
    """
    ps.dsp_orders = orders
    ps.dsp_cost = sum(money.D(o.total_cost) for o in orders)
    ps.dsp_impressions = sum(o.impressions for o in orders)
    ps.dsp_dpv = sum(o.total_dpv for o in orders)
    ps.dsp_atc = sum(o.total_atc for o in orders)
    ps.dsp_promoted_sales = sum(money.D(o.promoted_sales) for o in orders)
    ps.dsp_total_sales = sum(money.D(o.total_sales) for o in orders)
    ps.dsp_purchases = sum(o.total_units_sold for o in orders)  # 历史邮件「总销量」口径 = Total units sold
    ps.dsp_roas = float(ps.dsp_total_sales / ps.dsp_cost) if ps.dsp_cost else 0.0

    cons_orders = [o for o in orders if o.strategy == "Consideration"]
    conv_orders = [o for o in orders if o.strategy == "Conversion"]

    ps.consideration = _strategy_summary("Consideration", cons_orders)
    ps.conversion = _strategy_summary("Conversion", conv_orders)

    # 返点 / 付款金额口径：
    #   billing_base_cost = q2(dsp_cost)
    #   rebate_raw = billing_base_cost × rebate_rate
    #   payment_raw = billing_base_cost − rebate_raw
    #   rebate_amount = q2(rebate_raw)
    #   payment_amount = q2(payment_raw)
    ps.rebate_rate = money.D(rebate_rate)
    ps.billing_base_cost = money.q2(ps.dsp_cost)              # 49025.52
    ps.rebate_raw = ps.billing_base_cost * ps.rebate_rate     # 3922.0416
    ps.payment_raw = ps.billing_base_cost - ps.rebate_raw     # 45103.4784
    ps.rebate_amount = money.q2(ps.rebate_raw)                # 3922.04
    ps.payment_amount = money.q2(ps.payment_raw)              # 45103.48

    # DSP + 站内（站内部分由 promoted_ads_processor 汇总后调用 refresh_combined）
    ps.combined_cost = ps.dsp_cost
    ps.combined_sales = ps.dsp_total_sales


def _strategy_summary(strategy: str, orders: List[DSPOrder]) -> StrategySummary:
    s = StrategySummary(strategy=strategy)
    s.impressions = sum(o.impressions for o in orders)
    s.dpv = sum(o.total_dpv for o in orders)
    s.atc = sum(o.total_atc for o in orders)
    s.promoted_sales = sum(money.D(o.promoted_sales) for o in orders)
    s.total_cost = sum(money.D(o.total_cost) for o in orders)
    s.total_purchases = sum(o.total_purchases for o in orders)
    s.total_units_sold = sum(o.total_units_sold for o in orders)
    s.total_sales = sum(money.D(o.total_sales) for o in orders)
    return s


def refresh_combined(ps: PaymentSummary) -> None:
    """站内数据就绪后，刷新 DSP+站内 综合指标。"""
    ps.combined_cost = ps.dsp_cost + ps.onsite_cost
    ps.combined_sales = ps.dsp_total_sales + ps.onsite_sales
    ps.combined_roas = float(ps.combined_sales / ps.combined_cost) if ps.combined_cost else 0.0

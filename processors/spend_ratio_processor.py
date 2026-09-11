"""亚马逊美国DSP店铺花费占比处理。

店铺名称 / 固定比例**来自模板自身**，程序不新增/删除/改名/改比例，
也不根据 Order Summary 推导店铺。只修改：申请费用总额、费用月份、申请人、
各店铺 DSP费用（= 申请费用 × 模板固定比例）。

旧版 .xls：优先用 Windows Excel COM 原格式编辑（保留 .xls），
COM 不可用时回退 .xls → .xlsx 并明确告警，绝不静默转换。
"""
from __future__ import annotations

import datetime
import os
from decimal import ROUND_HALF_UP, Decimal
from typing import Dict, List, Optional, Tuple

import openpyxl

from models.payment_summary import PaymentSummary
from rules import amount_fields, month_rules
from utils import excel_com, money, xls_utils


def process_spend_ratio(template_path: str, ps: PaymentSummary, output_dir: str,
                        applicant: str = "余曼妮",
                        output_name: str = None) -> Tuple[str, List[str]]:
    logs: List[str] = []
    is_xls = template_path.lower().endswith(".xls") and not template_path.lower().endswith(".xlsx")

    if is_xls:
        # 读结构（店铺名 + 比例）
        import xlrd
        rb = xlrd.open_workbook(template_path, formatting_info=True)
        ws = rb.sheet_by_index(0)
        structure = _read_structure_xlrd(ws)
        _ensure_structure(structure)

        if excel_com.com_available():
            if output_name is None:
                output_name = f"亚马逊美国DSP店铺花费占比-{ps.month}_已处理.xls"
            out_path = os.path.join(output_dir, output_name)
            updates = _compute_updates(structure, ps, applicant)
            if excel_com.edit_xls_via_com(template_path, out_path, updates):
                _validate_preserved_xls(template_path, out_path, structure, logs)
                logs.append("✅ 店铺花费占比已生成（Excel COM 保留 .xls 格式）")
                return out_path, logs
            logs.append("⚠ Excel COM 编辑 .xls 失败，回退转换为 .xlsx")

        # 回退：转 .xlsx（明确告警）
        tmp_xlsx = os.path.join(output_dir, "_店铺占比模板_转换.xlsx")
        xls_utils.convert_xls_to_xlsx(template_path, tmp_xlsx)
        if output_name is None:
            output_name = f"亚马逊美国DSP店铺花费占比-{ps.month}_已处理.xlsx"
        out_path = _process_openpyxl(tmp_xlsx, ps, applicant, output_dir, output_name)
        logs.append("⚠ 当前系统无法原格式编辑 .xls，已转换为 .xlsx，无法保证旧版Excel所有格式/对象100%保留")
        return out_path, logs

    # .xlsx
    if output_name is None:
        output_name = f"亚马逊美国DSP店铺花费占比-{ps.month}_已处理.xlsx"
    out_path = _process_openpyxl(template_path, ps, applicant, output_dir, output_name)
    logs.append("✅ 店铺花费占比已生成")
    return out_path, logs


def _read_structure_xlrd(ws) -> dict:
    header_row = None
    for r in range(ws.nrows):
        if str(ws.cell_value(r, 0)).strip() == "核算维度":
            header_row = r
            break
    if header_row is None:
        raise ValueError("店铺花费占比模板未找到「核算维度」表头行")

    store_rows: Dict[str, int] = {}
    total_row: Optional[int] = None
    ratios: Dict[str, float] = {}
    for r in range(header_row + 1, ws.nrows):
        name = str(ws.cell_value(r, 0)).strip()
        if not name:
            continue
        if name == "合计":
            total_row = r + 1          # 转 1-based（COM Cells 用 1-based）
        else:
            store_rows[name] = r + 1   # 转 1-based
            ratios[name] = float(ws.cell_value(r, 2) or 0)
    return {"header_row": header_row, "store_rows": store_rows, "total_row": total_row, "ratios": ratios}


def _read_structure_openpyxl(ws) -> dict:
    header_row = None
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(row=r, column=1).value or "").strip() == "核算维度":
            header_row = r
            break
    if header_row is None:
        raise ValueError("店铺花费占比模板未找到「核算维度」表头行")
    store_rows: Dict[str, int] = {}
    total_row: Optional[int] = None
    ratios: Dict[str, float] = {}
    for r in range(header_row + 1, ws.max_row + 1):
        name = str(ws.cell(row=r, column=1).value or "").strip()
        if not name:
            continue
        if name == "合计":
            total_row = r
        else:
            store_rows[name] = r
            ratios[name] = float(ws.cell(row=r, column=3).value or 0)
    return {"header_row": header_row, "store_rows": store_rows, "total_row": total_row, "ratios": ratios}


def _ensure_structure(structure: dict) -> None:
    if not structure["store_rows"]:
        raise ValueError("店铺花费占比模板未识别到任何店铺")


def _compute_updates(structure: dict, ps: PaymentSummary, applicant: str) -> List[Tuple[int, int, object]]:
    """计算要写入的单元格 (row1based, col1based, value)。COM 写入用 float。"""
    apply_amount = money.q2(amount_fields.resolve_amount(ps, "spend_ratio.apply_amount"))
    first_store_row = min(structure["store_rows"].values())
    updates = []
    # D 申请费用总额（首个店铺行 + 合计行）
    updates.append((first_store_row, 4, float(apply_amount)))
    if structure["total_row"] is not None:
        updates.append((structure["total_row"], 4, float(apply_amount)))
    # E 申请人 / F 费用月份（首个店铺行）
    updates.append((first_store_row, 5, applicant))
    updates.append((first_store_row, 6, month_rules.excel_date_serial(ps.year, ps.month_num)))
    # B DSP费用 = 申请费用 × 模板固定比例
    for name, r in structure["store_rows"].items():
        ratio = money.D(structure["ratios"][name])
        updates.append((r, 2, float((apply_amount * ratio).quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP))))
    # 合计 B 列
    if structure["total_row"] is not None:
        updates.append((structure["total_row"], 2, float(apply_amount)))
    return updates


def _process_openpyxl(work_path: str, ps: PaymentSummary, applicant: str,
                      output_dir: str, output_name: str) -> str:
    wb = openpyxl.load_workbook(work_path)
    ws = wb[wb.sheetnames[0]]
    structure = _read_structure_openpyxl(ws)
    _ensure_structure(structure)
    before_stores = set(structure["store_rows"].keys())
    before_ratios = dict(structure["ratios"])

    apply_amount = money.q2(amount_fields.resolve_amount(ps, "spend_ratio.apply_amount"))
    first_store_row = min(structure["store_rows"].values())
    ws.cell(row=first_store_row, column=4, value=apply_amount)
    if structure["total_row"] is not None:
        ws.cell(row=structure["total_row"], column=4, value=apply_amount)
    ws.cell(row=first_store_row, column=5, value=applicant)
    ws.cell(row=first_store_row, column=6, value=month_rules.excel_date_serial(ps.year, ps.month_num))
    for name, r in structure["store_rows"].items():
        ratio = money.D(structure["ratios"][name])
        ws.cell(row=r, column=2, value=(apply_amount * ratio).quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP))
    if structure["total_row"] is not None:
        ws.cell(row=structure["total_row"], column=2, value=apply_amount)

    out_path = os.path.join(output_dir, output_name)
    wb.save(out_path)
    _validate_preserved_openpyxl(work_path, out_path, before_stores, before_ratios)
    return out_path


def _validate_preserved_openpyxl(template_path, out_path, before_stores, before_ratios) -> None:
    wb = openpyxl.load_workbook(out_path)
    ws = wb[wb.sheetnames[0]]
    after = _read_structure_openpyxl(ws)
    after_stores = set(after["store_rows"].keys())
    after_ratios = dict(after["ratios"])
    if before_stores != after_stores:
        raise ValueError(f"❌ 店铺花费占比模板结构被修改：店铺集合 {before_stores} → {after_stores}")
    for name in before_ratios:
        if name in after_ratios and abs(before_ratios[name] - after_ratios[name]) > 1e-9:
            raise ValueError(f"❌ 店铺花费占比比例被修改：{name} {before_ratios[name]} → {after_ratios[name]}")


def _validate_preserved_xls(template_path, out_path, structure, logs) -> None:
    """COM 路径后校验店铺/比例未被改动。"""
    import xlrd
    rb_out = xlrd.open_workbook(out_path)
    ws_out = rb_out.sheet_by_index(0)
    after = _read_structure_xlrd(ws_out)
    if set(structure["store_rows"].keys()) != set(after["store_rows"].keys()):
        raise ValueError("❌ 店铺花费占比模板结构被修改（店铺集合不一致）")
    for name, ratio in structure["ratios"].items():
        if abs(float(after["ratios"].get(name, -1)) - ratio) > 1e-9:
            raise ValueError(f"❌ 店铺花费占比比例被修改：{name}")
    logs.append("✅ 店铺/比例处理前后一致")

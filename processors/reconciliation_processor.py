"""DSP+SA 分析模板处理（可选）：把内部数据写入数据源 Sheet，刷新 DSP+SA Pivot。

数据源的真实来源：
  - DSP数据源 ← Order Summary（内部 List[DSPOrder]）
  - 已推广广告数据 数据源 ← 已推广-汇总（内部 List[PromotedAd]）

本模块只负责「若模板存在这些 Sheet 则写入留档/供 Pivot」，**不要求**模板必须
包含这些 Sheet。数据源缺失不报错，内部数据仍然可用。

DSP+SA Sheet 的策略表 / SA 表是 Pivot 输出、综合行是 GETPIVOTDATA 公式，不覆盖。
"""
from __future__ import annotations

import os
from typing import List, Optional, Tuple

import openpyxl

from models.payment_summary import DSPOrder, PaymentSummary, PromotedAd
from utils import excel_com, excel_utils


DSP_SRC_HEADERS = [
    "Order", "Strategy", "Total cost", "Impressions", "Total DPV", "Total DPVR",
    "Total ATC", "Total ATCR", "Promoted Sales", "Promoted ROAS", "Total purchases",
    "Total purchase rate", "Total units sold", "Total sales USD", "Total ROAS",
]
PROMO_SRC_HEADERS = [
    "店铺名称", "国家", "类型", "广告组合", "广告活动", "广告组", "广告组投放类型",
    "ASIN", "MSKU", "广告有效状态", "曝光量", "点击", "花费-本币",
    "销售额-本币", "广告订单", "直接成交销售额-本币",
]

DSP_SA_SHEET = ["DSP+SA", "SHEET-DSP+SA", "DSP+SA "]
DSP_SRC_SHEET = ["DSP数据源", "SHEET-DSP数据源"]
PROMO_SRC_SHEET = ["已推广广告数据 数据源", "SHEET-已推广广告数据 数据源"]


def _dsp_row(o: DSPOrder) -> list:
    return [o.order, o.strategy, o.total_cost, o.impressions, o.total_dpv, o.total_dpvr,
            o.total_atc, o.total_atcr, o.promoted_sales, o.promoted_roas, o.total_purchases,
            o.total_purchase_rate, o.total_units_sold, o.total_sales, o.total_roas]


def _promo_row(a: PromotedAd) -> list:
    return [a.store or "Amazon美国站哥贝尔店", a.country or "US", a.ad_type, a.portfolio,
            a.campaign, a.ad_group, a.target_type, a.asin, a.msku, a.status, a.impressions,
            a.clicks, a.cost, a.sales, a.orders, a.direct_sales]


def process_dsp_sa_template(template_path: str, ps: PaymentSummary, output_dir: str,
                            output_name: str = None, use_com: bool = True
                            ) -> Tuple[Optional[str], List[str], List[str]]:
    """处理 DSP+SA 分析模板（可选）。返回 (输出路径或 None, 日志, 被修改 Range)。

    - 若模板含 DSP数据源 / 已推广广告数据 数据源 Sheet → 写入（可选输出）。
    - 若不含 → 不报错，返回 None（内部数据仍然可用）。
    """
    logs: List[str] = []
    modified: List[str] = []

    wb = openpyxl.load_workbook(template_path, read_only=True)
    has_dsp_src = excel_utils.find_sheet(wb, DSP_SRC_SHEET) is not None
    has_promo_src = excel_utils.find_sheet(wb, PROMO_SRC_SHEET) is not None
    wb.close()

    # 处理前先校验模板自身表头是否完好，防止上传了「表头已被清空」的损坏模板
    _verify_template_header(template_path, has_dsp_src, has_promo_src)
    # 校验 Pivot 缓存字段名是否完好（防止「字段N」通用占位名导致 #NAME?）
    _verify_template_pivot(template_path)

    if not has_dsp_src and not has_promo_src:
        logs.append("ℹ 当前工作簿不包含《DSP数据源》/《已推广广告数据 数据源》Sheet，"
                    "内部 DSP/站内数据已由 Order Summary / 已推广-汇总 直接生成，跳过写入。")
        return None, logs, []

    if output_name is None:
        output_name = f"HUION-US-ENTITY-对账单-{ps.month_code}_已处理.xlsx"
    out_path = os.path.join(output_dir, output_name)

    dsp_rows = [_dsp_row(o) for o in ps.dsp_orders]
    promo_rows = [_promo_row(a) for a in ps.promoted_ads]

    if has_dsp_src and dsp_rows:
        modified.append(f"DSP数据源!A2:O{1 + len(dsp_rows)}")
    if has_promo_src and promo_rows:
        modified.append(f"已推广广告数据 数据源!A2:P{1 + len(promo_rows)}")

    # 记录处理前的 Pivot 结构（PivotTable/PivotCache 数量），处理后必须一致。
    pivot_before = excel_utils.pivot_signature(template_path)

    # 优先 COM（保留 Pivot/公式；只写数据源，不 RefreshAll、不重建 Pivot）
    if use_com and excel_com.com_available():
        ok = excel_com.write_reconciliation_via_com(
            template_path, out_path, dsp_rows, promo_rows,
            write_dsp=has_dsp_src, write_promo=has_promo_src)
        if ok:
            excel_utils.set_refresh_on_load_post_save(out_path)
            _verify_output(out_path, template_path, pivot_before, has_dsp_src, has_promo_src)
            logs.append("✅ 用 Excel COM 写入数据源（Pivot 打开文件时自动刷新）")
            if has_dsp_src:
                logs.append(f"✅ DSP数据源：{len(dsp_rows)} 条")
            else:
                logs.append("ℹ 模板不含《DSP数据源》Sheet，跳过")
            if has_promo_src:
                logs.append(f"✅ 已推广广告数据：{len(promo_rows)} 条")
            else:
                logs.append("ℹ 模板不含《已推广广告数据 数据源》Sheet，跳过")
            logs.append(f"✅ DSP+SA分析模板已生成：{output_name}")
            return out_path, logs, modified
        logs.append("⚠ Excel COM 写入失败，回退 openpyxl")

    # 回退 openpyxl（保留 Pivot 定义，仅更新数据源；Pivot 打开时由 Excel 刷新）
    logs.append("⚠ 未用 Excel COM：openpyxl 回退只更新数据源，DSP+SA 汇总表由 Excel 打开时刷新")
    wb = openpyxl.load_workbook(template_path)
    if has_dsp_src:
        _fill_dsp_source(excel_utils.find_sheet(wb, DSP_SRC_SHEET), dsp_rows)
    if has_promo_src:
        _fill_promo_source(excel_utils.find_sheet(wb, PROMO_SRC_SHEET), promo_rows)
    excel_utils.set_calc_mode_auto(wb)
    wb.save(out_path)
    # openpyxl 往返会打乱 Pivot 关系 ID，把模板的 Pivot 子树原样覆盖回输出，保证 Pivot 不被破坏。
    excel_utils.restore_pivot_parts(template_path, out_path)
    excel_utils.set_refresh_on_load_post_save(out_path)
    _verify_output(out_path, template_path, pivot_before, has_dsp_src, has_promo_src)
    logs.append(f"✅ DSP+SA分析模板已生成（openpyxl 回退）：{output_name}")
    return out_path, logs, modified


def _verify_template_header(template_path: str, has_dsp_src: bool, has_promo_src: bool) -> None:
    """处理前校验：模板自身的数据源 Sheet 第 1 行表头必须非空，否则判定为损坏模板。"""
    wb = openpyxl.load_workbook(template_path, read_only=True)
    try:
        for has_src, sheet_names, headers, label in (
            (has_dsp_src, DSP_SRC_SHEET, DSP_SRC_HEADERS, "《DSP数据源》"),
            (has_promo_src, PROMO_SRC_SHEET, PROMO_SRC_HEADERS, "《已推广广告数据 数据源》"),
        ):
            if not has_src:
                continue
            ws = excel_utils.find_sheet(wb, sheet_names)
            if ws is not None and _row_is_empty(ws, 1, len(headers)):
                actual = _read_row(ws, 1, len(headers))
                raise ValueError(
                    f"❌ 上传的对账单模板 {label} 第 1 行表头为空，模板可能已损坏（表头曾被清空）。\n"
                    f"   请重新上传完好的模板，表头应为：{headers}\n"
                    f"   当前读到的第 1 行：{actual}"
                )
    finally:
        wb.close()


def _verify_template_pivot(template_path: str) -> None:
    """处理前校验：模板 Pivot 缓存的字段名不能是「字段N」这类通用占位名。

    字段名被替换成 字段1/字段2/... 时，Excel 刷新 Pivot 会报 #NAME?，
    导致 DSP/站内 Pivot 只剩错位标签、缺 Strategy/类型 等分类。
    """
    import re
    import zipfile
    try:
        z = zipfile.ZipFile(template_path)
        names = z.namelist()
    except Exception:
        return  # 无法读取 zip，交由后续步骤报错

    generic_total = 0
    for name in names:
        if not re.match(r"xl/pivotCache/pivotCacheDefinition\d+\.xml$", name):
            continue
        xml = z.read(name).decode("utf-8", "replace")
        fields = re.findall(r'<cacheField[^>]*name="([^"]*)"', xml)
        generic_total += sum(1 for f in fields if re.match(r"^字段\d+$", f))
    z.close()

    if generic_total >= 3:
        raise ValueError(
            f"❌ 上传的对账单模板的 Pivot 缓存字段名已损坏（出现 {generic_total} 个「字段N」通用占位名）。\n"
            f"   这会导致 Pivot 刷新出现 #NAME? 错误，无法生成邮件截图。\n"
            f"   请重新上传完好的模板（字段名应为 Strategy/类型/曝光/总花费 等真实列名）。"
        )


def _verify_output(out_path: str, template_path: str, pivot_before: dict,
                   has_dsp_src: bool, has_promo_src: bool) -> None:
    """处理后校验：PivotTable/PivotCache 数量不得减少，数据源表头不得被清空。"""
    pivot_after = excel_utils.pivot_signature(out_path)
    for key, label in (("pivot_tables", "PivotTable"), ("pivot_caches", "PivotCache"),
                       ("pivot_cache_records", "PivotCacheRecords")):
        if pivot_after[key] < pivot_before[key]:
            raise ValueError(
                f"❌ {label} 数量减少：处理前 {pivot_before[key]} → 处理后 {pivot_after[key]}，处理失败。")

    wb = openpyxl.load_workbook(out_path, read_only=True)
    try:
        for has_src, sheet_names, headers, label in (
            (has_dsp_src, DSP_SRC_SHEET, DSP_SRC_HEADERS, "《DSP数据源》"),
            (has_promo_src, PROMO_SRC_SHEET, PROMO_SRC_HEADERS, "《已推广广告数据 数据源》"),
        ):
            if not has_src:
                continue
            ws = excel_utils.find_sheet(wb, sheet_names)
            if ws is None:
                raise ValueError(f"❌ 输出文件缺少 {label} Sheet，处理失败。\n输出文件：{out_path}")
            if _row_is_empty(ws, 1, len(headers)):
                actual = _read_row(ws, 1, len(headers))
                raise ValueError(
                    f"❌ {label} 第一行表头被清空，处理失败。\n"
                    f"   实际读到的第 1 行（前 {len(headers)} 列）：{actual}\n"
                    f"   预期表头：{headers}\n"
                    f"   输出文件：{out_path}\n"
                    f"   （模板表头是否完好？若模板本身表头为空，也会触发此报错）"
                )
    finally:
        wb.close()


def _read_row(ws, row: int, ncols: int) -> list:
    """读取某行前 ncols 列的值。"""
    return [ws.cell(row=row, column=c).value for c in range(1, ncols + 1)]


def _row_is_empty(ws, row: int, ncols: int) -> bool:
    """判断某行前 ncols 列是否全部为空。"""
    vals = _read_row(ws, row, ncols)
    return all(v is None or (isinstance(v, str) and not v.strip()) for v in vals)


def _fill_dsp_source(ws, dsp_rows: List[list]) -> None:
    excel_utils.clear_data_cells(ws, 2, max(ws.max_row, 2), 1, len(DSP_SRC_HEADERS))
    for i, vals in enumerate(dsp_rows):
        r = 2 + i
        for c, v in enumerate(vals, start=1):
            ws.cell(row=r, column=c, value=v)


def _fill_promo_source(ws, promo_rows: List[list]) -> None:
    excel_utils.clear_data_cells(ws, 2, max(ws.max_row, 2), 1, len(PROMO_SRC_HEADERS))
    for i, vals in enumerate(promo_rows):
        r = 2 + i
        for c, v in enumerate(vals, start=1):
            ws.cell(row=r, column=c, value=v)

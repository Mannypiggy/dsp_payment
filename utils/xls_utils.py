"""旧版 .xls 安全兼容：用 xlrd 读取 → openpyxl 写出 .xlsx。

原因：xlutils/xlwt 在 Python 3.14 下因 None 数字格式字符串回写崩溃，
故对 .xls 模板采用「读取内容 + 结构，转写为 .xlsx」的兼容方案，
保留单元格值、类型、合并单元格、列宽、数字格式（日期/百分比/货币）。
"""
from __future__ import annotations

import datetime
import os
from typing import List, Tuple

import openpyxl
import xlrd

# xlrd 内置格式索引 -> openpyxl 数字格式
_XLRD_BUILTIN = {
    0: None,             # General
    1: "0",
    2: "0.00",
    3: "#,##0",
    4: "#,##0.00",
    9: "0%",
    10: "0.00%",
    14: "yyyy-mm-dd",    # m/d/yy
    49: "@",             # text
}


def _fmt_key_to_str(rb: xlrd.Book, xf_index: int):
    try:
        xf = rb.xf_list[xf_index]
        fmt_key = xf.format_key
        if fmt_key in _XLRD_BUILTIN:
            return _XLRD_BUILTIN[fmt_key]
        fmt = rb.format_map.get(fmt_key)
        if fmt is not None:
            return fmt.format_str
    except Exception:
        return None
    return None


def convert_xls_to_xlsx(xls_path: str, xlsx_path: str) -> str:
    """把 .xls 转成 .xlsx，返回 xlsx 路径。"""
    rb = xlrd.open_workbook(xls_path, formatting_info=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for sheet in rb.sheets():
        ws = wb.create_sheet(title=sheet.name)
        # 合并单元格
        for (rlo, rhi, clo, chi) in sheet.merged_cells:
            ws.merge_cells(start_row=rlo + 1, start_column=clo + 1,
                           end_row=rhi, end_column=chi)
        # 值 + 数字格式
        for r in range(sheet.nrows):
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                ct = cell.ctype
                if ct == xlrd.XL_CELL_EMPTY or ct == xlrd.XL_CELL_BLANK:
                    continue
                oc = ws.cell(row=r + 1, column=c + 1)
                if ct == xlrd.XL_CELL_DATE:
                    oc.value = _date_to_datetime(cell.value)
                elif ct == xlrd.XL_CELL_NUMBER:
                    oc.value = float(cell.value)
                elif ct == xlrd.XL_CELL_BOOLEAN:
                    oc.value = bool(cell.value)
                elif ct == xlrd.XL_CELL_ERROR:
                    continue
                else:
                    oc.value = cell.value
                nf = _fmt_key_to_str(rb, sheet.cell_xf_index(r, c))
                if nf:
                    oc.number_format = nf
        # 列宽（xlrd 单位：字符数的 1/256；openpyxl 用字符宽）
        for c in range(sheet.ncols):
            colinfo = sheet.colinfo_map.get(c)
            if colinfo and colinfo.width and colinfo.width != 256 * 8:
                width = colinfo.width / 256.0
                ws.column_dimensions[openpyxl.utils.get_column_letter(c + 1)].width = width

    wb.save(xlsx_path)
    return xlsx_path


def _date_to_datetime(v) -> datetime.datetime:
    """xlrd 日期（浮点序列号）→ datetime。"""
    try:
        year, month, day, hour, minute, second = xlrd.xldate_as_tuple(v, 0)
        return datetime.datetime(year, month, day, hour, minute, second)
    except Exception:
        return v

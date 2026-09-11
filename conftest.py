"""pytest 共享 fixture：真实 2026-07 历史数据路径 + 构建 PaymentSummary 的辅助。"""
from __future__ import annotations

import os
import sys

import pytest

# 让 tests 内可直接 import 顶层包
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

JULY_BASE = r"C:\Users\41250\Desktop\2026年7月Huion-US对账单\2026年7月Huion-US对账单"
MEETING_PATH = r"C:\Users\41250\Desktop\2026年7月Huion-US对账单\2026年7月Huion-US对账单\测试\7月dsp会议记录.docx"


@pytest.fixture(scope="session")
def july_base():
    return JULY_BASE


@pytest.fixture(scope="session")
def july_order_summary():
    return os.path.join(JULY_BASE, "HUION_US_(Advertiser)_Order_Summary_20260810015543.xlsx")


@pytest.fixture(scope="session")
def july_reconciliation():
    return os.path.join(JULY_BASE, "HUION-US-ENTITY-对账单-自用.xlsx")


@pytest.fixture(scope="session")
def july_spend_detail():
    return os.path.join(JULY_BASE, "DSP花费明细-202607.xlsx")


@pytest.fixture(scope="session")
def july_spend_ratio():
    return os.path.join(JULY_BASE, "亚马逊美国DSP店铺花费占比-2026-07.xls")


@pytest.fixture(scope="session")
def july_meeting():
    return MEETING_PATH


@pytest.fixture(scope="session")
def july_invoice():
    return os.path.join(JULY_BASE, "Invoice_M001098_HUION GLOBAL (HK) LIMITED.pdf")


@pytest.fixture(scope="session")
def july_statement():
    """付款对账单（对账清单-待付款，用于核对最终付款金额）。"""
    return os.path.join(JULY_BASE, "2026年7月对账单SparkX&Huion-US.xlsx")


@pytest.fixture(scope="session")
def july_promoted_extract(tmp_path_factory):
    """从对账单提取「已推广广告数据 数据源」为独立 xlsx（模拟已处理后的源）。"""
    import openpyxl
    src = os.path.join(JULY_BASE, "HUION-US-ENTITY-对账单-自用.xlsx")
    wb = openpyxl.load_workbook(src, data_only=True)
    ws = wb["已推广广告数据 数据源"]
    out = tmp_path_factory.mktemp("data") / "已推广数据源_提取.xlsx"
    w2 = openpyxl.Workbook()
    ws2 = w2.active
    for row in ws.iter_rows(values_only=True):
        ws2.append(list(row))
    w2.save(str(out))
    return str(out)


def build_july_ps():
    """用真实 2026-07 数据构建完整 PaymentSummary（供邮件/返点/KPI 测试）。"""
    from utils.config_loader import load_config
    from rules import month_rules
    from processors import (order_summary_processor, promoted_ads_processor,
                            kpi_processor)

    config = load_config()
    ps = month_rules.parse_month("2026-07", peak_months=config["month"]["peak_months"])
    orders = order_summary_processor.parse_order_summary(
        os.path.join(JULY_BASE, "HUION_US_(Advertiser)_Order_Summary_20260810015543.xlsx"),
        store_default=config["store"]["default_store"],
    )
    order_summary_processor.aggregate(ps, orders, rebate_rate=config["payment"]["rebate_rate"])

    # 提取已推广数据
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(JULY_BASE, "HUION-US-ENTITY-对账单-自用.xlsx"), data_only=True)
    ws = wb["已推广广告数据 数据源"]
    from models.payment_summary import PromotedAd
    from rules import promoted_ads_mapping
    ads = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
        if row[2] is None:
            continue
        ad = PromotedAd(
            ad_type=str(row[2]).strip(),
            portfolio=str(row[3] or ""),
            campaign=str(row[4] or ""),
            ad_group=str(row[5] or ""),
            asin=str(row[7] or ""),
            msku=str(row[8] or ""),
            status=str(row[9] or ""),
            impressions=float(row[10] or 0),
            clicks=float(row[11] or 0),
            cost=float(row[12] or 0),
            sales=float(row[13] or 0),
            orders=float(row[14] or 0),
            direct_sales=float(row[15] or 0),
        )
        ads.append(ad)
    promoted_ads_processor.aggregate_onsite(ps, ads)
    kpi_processor.compute_and_reconcile(ps, targets=config["kpi"], template_values=None)
    ps.payer = config["payment"]["payer"]["name"]
    ps.payee = config["payment"]["payee"]
    return ps
